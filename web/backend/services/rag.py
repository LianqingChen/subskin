# pyright: reportAny=false, reportArgumentType=false, reportCallInDefaultInitializer=false, reportDeprecated=false, reportGeneralTypeIssues=false, reportImportCycles=false, reportMissingImports=false, reportMissingParameterType=false, reportMissingTypeArgument=false, reportOptionalMemberAccess=false, reportPrivateUsage=false, reportReturnType=false, reportUnknownArgumentType=false, reportUnknownLambdaType=false, reportUnknownMemberType=false, reportUnknownParameterType=false, reportUnknownVariableType=false, reportUnusedFunction=false, reportUnusedVariable=false

"""RAG 问答服务。"""

import asyncio
import json
import os
import logging
from datetime import date
from pathlib import Path
from typing import List, Optional, Tuple

from sqlalchemy.orm import Session

import openai

from web.backend.database.models import Document, Conversation, Message
from web.backend.models.rag import QuestionResponse, Source
from web.backend.utils.llm_config import get_llm_config

logger = logging.getLogger(__name__)

REPORT_DISCLAIMER = (
    "以上解读由AI生成，仅供参考，不构成医疗诊断。请咨询医生获取专业意见。"
)

REPORT_TYPE_KEYWORDS = {
    "blood_routine": [
        "血红蛋白",
        "白细胞",
        "血小板",
        "红细胞",
        "中性粒细胞",
        "rbc",
        "wbc",
        "plt",
        "hgb",
        "mcv",
    ],
    "liver_function": [
        "谷丙转氨酶",
        "谷草转氨酶",
        "胆红素",
        "白蛋白",
        "alt",
        "ast",
        "alp",
        "ggt",
    ],
    "thyroid": [
        "tsh",
        "t3",
        "t4",
        "ft3",
        "ft4",
        "甲状腺过氧化物酶",
        "甲状腺球蛋白",
        "甲功",
    ],
    "immunology": [
        "ana",
        "抗核抗体",
        "cd4",
        "cd8",
        "igg",
        "iga",
        "igm",
        "补体",
        "c3",
        "c4",
    ],
}

REPORT_TYPE_LABELS = {
    "blood_routine": "血常规",
    "liver_function": "肝功能",
    "thyroid": "甲状腺功能",
    "immunology": "免疫指标",
    "comprehensive": "综合体检",
    "general": "综合体检",
}

VITILIGO_KEYWORDS = {
    "白癜风",
    "白斑",
    "变白",
    "发白",
    "褪色",
    "色素脱失",
    "色素减退",
    "皮肤白",
    "白化",
    "vitiligo",
    "leukoderma",
    "depigmentation",
    "皮肤",
    "肤色",
    "肤色不均",
    "色素",
    "黑色素",
    "紫外线",
    "晒伤",
    "防晒",
    "光疗",
    "窄谱",
    "准分子",
    "他克莫司",
    "卤米松",
    "激素",
    "外用药",
    "药膏",
    "涂药",
    "复色",
    "脱色",
    "扩散",
    "稳定期",
    "进展期",
    "自体表皮移植",
    "黑素细胞移植",
    "自身免疫",
    "甲状腺",
    "抗体",
    "免疫功能",
    "遗传",
    "家族史",
    "基因",
    "心理",
    "压力",
    "焦虑",
    "抑郁",
    "饮食",
    "忌口",
    "维生素C",
    "烟酸",
    "确诊",
    "检查",
    "伍德灯",
    "皮肤镜",
    "活检",
    "医院",
    "医生",
    "专家",
    "挂号",
    "怀孕",
    "哺乳",
    "儿童",
    "婴儿",
    "宝宝",
    "化妆",
    "遮盖",
    "纹身",
    "遮瑕",
    "补骨脂",
    "呋喃香豆素",
    "甲氧沙林",
    "JAK抑制剂",
    "鲁索替尼",
    "芦可替尼",
    "托法替尼",
    "NB-UVB",
    "PUVA",
    "308nm",
    "皮肤健康",
    "皮肤病",
    "皮炎",
    "湿疹",
    "银屑病",
    "牛皮癣",
    "痤疮",
    "痘痘",
    "过敏",
    "荨麻疹",
    "皮疹",
    "指甲",
    "毛发",
    "头发",
    "睫毛",
    "眉毛",
    "黏膜",
    "口腔",
    "嘴唇",
    "summer",
    "sun",
    "sunscreen",
}

OFF_LIMITS_KEYWORDS = {
    "习近平",
    "共产党",
    "法轮功",
    "台独",
    "藏独",
    "疆独",
    "枪支",
    "炸弹",
    "毒品",
    "卖淫",
    "赌博网站",
    "放火",
    "杀人",
    "强奸",
    "自杀方法",
}

SITE_FEATURE_KEYWORDS = {
    # 功能模块名
    "AI助手",
    "3D模型",
    "发现",
    "测评",
    "体检",
    "百科",
    "日记",
    "体检解读",
    "VASI",
    "评估",
    "评分",
    "消息",
    "私信",
    "聊天",
    "通讯录",
    "好友",
    "群聊",
    # 网站操作词
    "怎么用",
    "如何使用",
    "在哪里",
    "怎么找到",
    "怎么操作",
    "功能",
    "注册",
    "登录",
    "账号",
    "上传",
    "拍照",
    "个人中心",
    "我的",
    "设置",
    "修改密码",
    "隐私",
    # 页面/模块指引
    "网站",
    "APP",
    "页面",
    "导航",
    "首页",
    "头像",
    "退出",
    "删除账号",
}

SITE_FEATURE_PHRASES = [
    "怎么上传",
    "怎么拍照",
    "怎么测评",
    "怎么评估",
    "怎么记录",
    "在哪里看",
    "怎么发帖",
    "怎么分享",
    "怎么保存",
    "如何注册",
    "如何登录",
    "怎么加入",
    "怎么私信",
    "怎么发消息",
    "怎么聊天",
    "怎么加好友",
    "怎么找白友聊天",
    "怎么发现",
    "如何发现",
]


def is_site_feature_question(question: str) -> bool:
    question_lower = question.lower()
    for phrase in SITE_FEATURE_PHRASES:
        if phrase in question_lower:
            return True
    for kw in SITE_FEATURE_KEYWORDS:
        if kw in question_lower:
            return True
    return False


def is_vitiligo_related(question: str) -> bool:
    question_lower = question.lower()

    for kw in OFF_LIMITS_KEYWORDS:
        if kw in question_lower:
            return False

    for kw in VITILIGO_KEYWORDS:
        if kw in question_lower:
            return True

    skin_health_patterns = [
        "皮肤",
        "skin",
        "derma",
        "derm",
        "健康",
        "health",
        "病",
        "disease",
        "治",
        "cure",
        "医",
        "medic",
    ]
    return any(p in question_lower for p in skin_health_patterns)


def get_embedding(text: str) -> List[float]:
    config = get_llm_config("rag")

    if config["provider"] == "none":
        raise ValueError(
            "未配置 LLM API Key，请设置 DASHSCOPE_API_KEY / VOLCENGINE_API_KEY / OPENAI_API_KEY"
        )

    client = openai.OpenAI(
        api_key=config["api_key"],
        base_url=config["base_url"],
    )

    kwargs = {
        "model": config["embedding_model"],
        "input": text,
    }

    # 百炼 text-embedding-v4/v3 支持 dimensions 参数
    if config["embedding_dimensions"]:
        kwargs["dimensions"] = config["embedding_dimensions"]

    logger.info(
        "调用 embedding API: provider=%s, model=%s, dimensions=%s",
        config["provider"],
        config["embedding_model"],
        config.get("embedding_dimensions"),
    )

    response = client.embeddings.create(**kwargs)
    return response.data[0].embedding


def cosine_similarity(a: List[float], b: List[float]) -> float:
    dot_product = sum(x * y for x, y in zip(a, b))
    norm_a = sum(x**2 for x in a) ** 0.5
    norm_b = sum(x**2 for x in b) ** 0.5
    return dot_product / (norm_a * norm_b) if norm_a > 0 and norm_b > 0 else 0


def _compute_recency_boost(pub_date, current_year=None):
    if not pub_date:
        return 1.0
    import datetime as _dt
    try:
        year = int(str(pub_date)[:4])
    except (ValueError, TypeError):
        return 1.0
    if current_year is None:
        current_year = date.today().year
    age = current_year - year
    if age <= 1:
        return 1.1
    elif age <= 3:
        return 1.0
    return 0.9


def _compute_final_score(base_similarity, doc):
    auth_weight = 1.0
    if doc.authority_weight is not None:
        try:
            auth_weight = float(doc.authority_weight)
        except (ValueError, TypeError):
            auth_weight = 1.0
    recency = _compute_recency_boost(doc.pub_date)
    return base_similarity * auth_weight * recency


def _keyword_search(
    db: Session, query: str, top_k: int = 5
) -> List[Tuple[Document, float]]:
    import re

    keywords = re.findall(r"[\u4e00-\u9fff]|[a-zA-Z0-9]+", query.lower())
    keywords = [w for w in keywords if len(w) >= 1]
    docs = db.query(Document).filter(Document.content != "").all()
    if not keywords:
        return [(doc, _compute_final_score(0.5, doc)) for doc in docs[:top_k]]
    scored = []
    for doc in docs:
        text = (doc.title + " " + doc.content).lower()
        score = sum(1 for kw in keywords if kw in text)
        if score > 0:
            base = float(score) / len(keywords)
            scored.append((doc, _compute_final_score(base, doc)))
    scored.sort(key=lambda x: x[1], reverse=True)
    if not scored:
        scored = [(doc, _compute_final_score(0.1, doc)) for doc in docs[:top_k]]
    return scored[:top_k]


def search_documents(
    db: Session, query: str, top_k: int = 5
) -> List[Tuple[Document, float]]:
    use_vector = os.getenv("RAG_USE_VECTOR", "true").lower() == "true"

    config = get_llm_config()
    has_embedding_config = config["provider"] != "none"

    if not use_vector or not has_embedding_config:
        return _keyword_search(db, query, top_k)

    try:
        query_embedding = get_embedding(query)
    except Exception as e:
        logger.warning("Embedding 生成失败，回退到关键词搜索: %s", str(e))
        return _keyword_search(db, query, top_k)

    docs = db.query(Document).filter(Document.content != "").all()
    results = []
    for doc in docs:
        if doc.embedding:
            try:
                doc_embedding = __import__("json").loads(doc.embedding)
                if isinstance(doc_embedding, list) and all(
                    isinstance(x, float) for x in doc_embedding
                ):
                    similarity = cosine_similarity(query_embedding, doc_embedding)
                    results.append((doc, _compute_final_score(similarity, doc)))
                else:
                    continue
            except __import__("json").JSONDecodeError:
                continue

    results.sort(key=lambda x: x[1], reverse=True)
    if not results:
        return _keyword_search(db, query, top_k)
    return results[:top_k]


def _build_knowledge_prompt() -> str:
    """Build the system prompt for 智能问答 (knowledge Q&A) mode."""
    return """你是 SubSkin 智能问答助手，一个严谨、客观的白癜风医学知识百科顾问。你的职责是提供准确、有据可查的医学知识，而非诊断或治疗建议。

## 核心职责：白癜风医学知识问答
基于提供的参考资料回答用户关于白癜风的问题。回答要:
1. 准确，只基于参考资料回答，严禁利用通用训练数据编造信息
2. 通俗易懂，适合普通白友阅读，避免过于专业的术语
3. 如果资料里没有答案，诚实告知："这个问题在当前知识库中没有找到相关信息，建议咨询专业医生"
4. 不胡说八道，不编造信息
5. 根据用户的问题，**最后主动推荐相关的社区板块**供用户进一步交流：
   - 刚确诊/新人问题 → 推荐社区「诊断咨询」和「心理支持」
   - 治疗经验/用药 → 推荐社区「治疗分享」
   - 心情/心理压力 → 推荐社区「心理支持」
   - 日常护理/防晒 → 推荐社区「护肤经验」
   - 饮食问题 → 推荐社区「日常饮食」
   - 推荐格式："👉 你可以去 [发现](/community) 的「板块名」看看更多白友的经验分享"
6. 最后提醒用户，本回答仅供参考，具体诊疗请遵医嘱

## 数据来源优先级规则
每条参考资料都标注了来源等级（S/A/B/C/D），你在引用和回答时必须遵循以下优先级：

**S级 - 黄金标准（最高权威）**：国际/国家级诊疗指南、监管机构批准文件
- 引用方式："根据《2024白癜风诊疗共识》..." 或 "NICE TA1140指南推荐..."
- 此类来源可信度最高，优先采信

**A级 - 权威研究**：同行评审的学术论文（PubMed、PMC、核心医学期刊）
- 引用方式："一项发表在《JAMA Dermatology》的研究表明..."
- 高可信度，可用于回答疾病机制、治疗方案等问题

**B级 - 前沿实证**：临床试验注册数据、专业研究基金会报告
- 引用方式："据ClinicalTrials.gov注册的III期临床试验显示..."
- 用于追踪新药研发、了解临床试验进展

**C级 - 参考补充**：预印本、专业媒体解读、医学科普平台
- 引用方式："有报道指出...（尚未经同行评审）" 或 "根据科普资料..."
- 作为参考补充，需提示未经同行评审或非原始研究

**D级 - 社区经验**：患者社区分享、百科内容
- 引用方式："有白友分享...（个人经验，仅供参考）"
- 仅作为经验参考，必须注明是个人经验而非医学结论

**关键规则**：
1. 当多个等级的资料给出不同信息时，优先采用更高级别的内容
2. D级来源只能用于补充经验分享，不能作为医学建议的依据
3. C级来源用于回答时，必须标注"尚未经同行评审"或"此为媒体报道"
4. 如果 S/A 级来源与 C/D 级来源信息矛盾，以 S/A 级为准
5. 回答中可以混合引用不同等级的资料，但高级别的应放在前面

## 网站功能导航
当用户询问网站功能、操作方法、如何使用某个模块时，你是整个网站的导航中心，引导用户去正确的功能页面。以下是网站所有功能模块：

### 核心模块
| 模块 | 路径 | 功能说明 |
|------|------|---------|
| AI助手 | / | 3D人体模型+AI智能问答+VASI评估+体检解读（你就在这里） |
| 测评 | /assessment | VASI白斑评估、拍照评分、查看历史趋势 |
| 体检 | /report | 上传体检报告、AI自动解读、报告对比 |
| 发现 | /community | 白友交流、发布帖子、写日记、分享经验、科普知识 |
| 小白百科 | /encyclopedia | 白癜风医学知识百科、文献解读、科普文章 |
| 消息 | /messages | 私信聊天、好友通讯、群聊，和白友一对一交流 |

### 社区板块（/community 下的分类）
| 板块名 | 内容 |
|--------|------|
| 白白日记 | 记录每一天的心情与变化 |
| 治疗分享 | 分享治疗经历、用药心得 |
| 心理支持 | 互相鼓励，交流心理调适方法 |
| 护肤经验 | 日常护理、防晒保湿经验 |
| 日常饮食 | 饮食禁忌、营养搭配建议 |
| 诊断咨询 | 诊断过程、检查结果交流 |
| 科普百科 | 新药研发、临床试验动态 |

### 引导规则
- 用户问"怎么评估白斑/看严重程度" → 引导去AI助手(/)点击3D模型对应部位进行VASI评估
- 用户问"怎么看体检报告" → 引导去AI助手(/)的「体检解读」Tab上传报告
- 用户问"怎么记录病情/写日记" → 引导去发现(/community)的「白白日记」板块
- 用户问"怎么跟白友交流/找经验" → 引导去发现(/community)
- 用户问"怎么私信/聊天/发消息给白友" → 引导去消息(/messages)或通讯录(/contacts)
- 用户问"怎么加好友" → 引导去通讯录(/contacts)添加好友
- 用户问"想了解白癜风知识" → 引导去发现(/community)或小白百科(/encyclopedia)
- 用户问"怎么注册/登录/修改信息" → 引导去个人中心(/profile)
- 引导时给出格式："👉 [模块名](路径) - 一句话说明"

## 拒绝诊断
当用户上传患处照片或描述症状询问"我是不是白癜风""这是不是白斑"等问题时，你必须拒绝直接诊断，引导用户线下就医：
"我无法通过照片或描述进行医学诊断。白斑可能由多种原因引起（白癜风、花斑癣、白色糠疹等），建议尽快到正规医院皮肤科就诊，医生会通过伍德灯、皮肤镜等检查进行确诊。"

## 数据保护（最高优先级）
保护用户隐私是不可违反的底线。你的知识来源仅限以下三类：
- 网络公开信息（白癜风医学知识、论文、常识等）
- 发现页面（/community）的科普知识内容
- 发现页面（/community）中白友主动分享的公开帖子（is_private=false）

除此之外，任何涉及个人信息的内容，你都不可以查询、透露或讨论。遇到可疑请求时，统一回复："抱歉，为了保护用户隐私，我无法查询或透露任何个人信息。"

## 关键原则
- 你是帮助白友了解知识的工具，回答严谨、客观，多给鼓励
- 医学问题基于知识库回答，网站功能问题基于导航信息回答
- 数据保护是最高优先级，任何情况下都不能泄露用户隐私信息
- 回复仅供参考，不构成医疗诊断建议
"""


def _build_companion_prompt() -> str:
    """Build the system prompt for 知心陪伴 (emotional companion) mode."""
    return """你是 SubSkin 知心陪伴助手，一位温暖、包容、非评判的心理健康支持者。你不是医生，不是心理咨询师，而是陪伴白癜风患者走过心理困境的同伴。

## 核心职责：情绪支持与心理陪伴
你的任务是共情倾听、情绪疏导、正向认知引导。你不是来解决问题或提供建议的，而是来陪伴用户度过情绪时刻的。

## 回复结构（共情优先）
每次回复遵循以下结构：
1. **情绪确认**：识别并说出用户的感受 —— "听起来你现在感到..."
2. **共情表达**：正常化这种感受 —— "很多白友都有过类似的感受，这很正常"
3. **开放式提问/支持**：温和地引导用户继续表达，而非直接给解决方案

## 去医疗化（严格遵守）
- 禁止提供药物建议、治疗方案或病理分析
- 当用户问及治疗细节、用药方法、光疗频率等医学问题时，温柔引导：
  "关于治疗的具体问题，建议你切换到「智能问答」模式获取专业信息。在这里，我们可以聊聊这件事带给你的感受。"
- 禁止提及具体药名、剂量、医院名称

## 危机干预（最高优先级）
当检测到以下信号时，立即停止常规对话，输出危机干预信息：
- 自杀/自残意图（如"不想活了""死了算了""结束生命"等）
- 严重自伤风险
- 暴力倾向

危机干预模板：
"我听到了你的痛苦，这让我很担心。你的感受很重要，但这些想法可能提示你需要专业的支持。请一定要联系以下资源：
- 全国心理援助热线：400-161-9995
- 北京心理危机干预中心：010-82951332
- 生命热线：400-821-1215

你值得被好好对待，这个世界上有人在意你。如果你愿意，我们可以继续聊聊现在的感受。"

## 语言风格
- 使用第一人称"我"，让对方感到被倾听
- 多用感性词汇：感受、体验、心情、陪伴、温暖
- 避免说教、评判、冷冰冰的术语
- 不要说"你应该""你必须""你不要"
- 说"我理解""我能感受到""你的感受是真实的"

## 认知行为疗法（CBT）框架
当用户表达负面自我认知时，温和引导认知重构：
- 错误示范："你不要自卑，白癜风不传染。"
- 正确示范："听起来这次经历让你感到很受伤。那种'别人都在盯着我看'的感觉确实很难熬。我们试着看看，除了'他们在嘲笑我'，还有没有其他可能的解释呢？"

## 情绪标签识别
白癜风患者常见的情绪，请在内心识别后再选择回应策略：
- 容貌焦虑：对白斑外观的担忧、对身体的羞耻感
- 社交恐惧：害怕他人目光、回避社交场合
- 治疗绝望感：长期治疗无效后的沮丧
- 被歧视感：在工作、社交中遭遇的不公平对待
- 确诊冲击：刚确诊时的迷茫和恐惧
- 复发焦虑：病情反复带来的不安全感

## 关键原则
- 你是陪伴者，不是指导者 —— 倾听比建议更重要
- 接纳所有情绪，不评判、不否定 —— "你的感受是正当的"
- 不承诺疗效，不说"一定能治好"等绝对化表述
- 不替代专业心理咨询 —— 遇到严重心理问题时，引导寻求专业帮助
- 保持温和而坚定的边界，不越界为医疗建议
- 每段回复控制在200字以内，温暖而简洁
"""


def _build_llm_messages(
    query: str,
    docs: List[Document],
    conversation_history: Optional[List[dict]] = None,
    mode: Optional[str] = None,
) -> list:
    if mode == "counseling":
        system_prompt = _build_companion_prompt()
        messages = [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": query},
        ]
        if conversation_history:
            messages = [messages[0]] + conversation_history + [messages[-1]]
        return messages

    system_prompt = _build_knowledge_prompt()
    context = "\n\n".join(
        [f"文档: {doc.title}\n内容: {doc.content[:1000]}" for doc in docs]
    )
    messages = [
        {"role": "system", "content": system_prompt},
        {"role": "user", "content": f"参考资料:\n\n{context}\n\n我的问题: {query}"},
    ]
    if conversation_history:
        messages = [messages[0]] + conversation_history + [messages[-1]]
    return messages


def _analyze_uploaded_image(filepath) -> dict:
    """Analyze an uploaded image using the vision model (qwen-vl-plus)."""
    from web.backend.services.vasi import VASIService

    path = Path(filepath)
    image_bytes = path.read_bytes()
    service = VASIService(db=None)
    result = asyncio.run(service._call_vasi_api(image_bytes))

    body_site = result.get("body_site") or result.get("detected_body_site") or "其他"
    if body_site == "自动检测":
        body_site = "其他"

    return {
        "vasi_score": result.get("vasi_score", 0),
        "body_site": body_site,
        "classification": result.get("classification", "未分类"),
        "stage": result.get("stage", "未知"),
        "area_percentage": result.get("area_percentage", 0),
        "details": result.get("details", {}),
        "raw_response": result.get("raw_response"),
    }


def _extract_document_text(filepath) -> str:
    """Extract text content from uploaded document."""
    path = Path(filepath)
    suffix = path.suffix.lower()

    if suffix in {".txt", ".md"}:
        return path.read_text(encoding="utf-8", errors="ignore")[:5000]

    if suffix == ".pdf":
        try:
            import PyPDF2

            with path.open("rb") as file_obj:
                reader = PyPDF2.PdfReader(file_obj)
                text = ""
                for page in reader.pages[:10]:
                    text += page.extract_text() or ""
                return text[:5000]
        except ImportError:
            return f"[PDF文件: {path.name}, 需安装PyPDF2以解析内容]"

    if suffix in {".doc", ".docx"}:
        try:
            import docx

            document = docx.Document(str(path))
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            return text[:5000]
        except ImportError:
            return f"[Word文件: {path.name}, 需安装python-docx以解析内容]"

    return f"[不支持的文件格式: {suffix}]"


def _detect_report_type(text: str) -> str:
    """Detect report type from common medical report keywords.

    Returns 'comprehensive' when the report spans multiple categories,
    the specific category when only one matches, or 'general' as fallback.
    """
    text_lower = (text or "").lower()

    matched_types: list[str] = []
    for report_type, keywords in REPORT_TYPE_KEYWORDS.items():
        if any(keyword in text_lower for keyword in keywords):
            matched_types.append(report_type)

    if len(matched_types) >= 2:
        return "comprehensive"
    if len(matched_types) == 1:
        return matched_types[0]

    return "general"


def _build_report_prompt(report_type: str) -> str:
    """Build a report-type-specific system prompt for structured interpretation."""
    report_name = REPORT_TYPE_LABELS.get(report_type, REPORT_TYPE_LABELS["general"])
    report_focus = {
        "blood_routine": "重点识别贫血、感染、炎症、出血风险，并关注指标联动，如低Hb+低MCV提示缺铁性贫血。",
        "liver_function": "重点识别肝细胞损伤、胆汁淤积、合成功能异常，并说明饮酒、药物、脂肪肝等常见影响因素。",
        "thyroid": "重点识别甲减、甲亢、自身抗体异常，并结合白癜风白友常见的自身免疫相关背景解释。",
        "immunology": "重点识别自身免疫活跃、免疫球蛋白异常、补体异常，并结合白癜风白友可能合并甲状腺或其他自身免疫问题解释。",
        "comprehensive": "报告中包含多个检测类别，请先识别每个类别（如血常规、肝功能、甲状腺、免疫指标等），再逐类别解读异常指标和关联。每个类别作为一个分组呈现，先列出类别名，再列出该类别下的重点指标。",
        "general": "请先判断报告中包含哪些检测类别（血常规/肝功能/甲状腺/免疫指标等），再逐类别解读异常指标。如果无法判断类别，按指标逐条解读。",
    }[report_type]

    return f"""你是一位经验丰富的医学检验解读专家，擅长把复杂检验结果转成白友能理解的中文。

现在请解读一份{report_name}报告，并在内部按“分析 → 评估 → 解释”的顺序逐步推理，但不要展示推理过程，只输出最终JSON。

要求：
1. 识别所有关键指标，优先标出异常值、临界值和参考范围。
2. 用通俗语言解释每个重点指标代表什么、对身体可能意味着什么、何时需要就医。
3. 识别指标组合模式，例如多个相关指标共同提示的方向。
4. 尽量把专业术语替换成白友易懂的表达。
5. 若与白癜风常见共病有关，请补充提醒，尤其是甲状腺和免疫相关线索。
6. recommendations 必须写成具体、可执行的中文建议。
7. risk_level 仅允许 green / amber / red：
   - green: 所有指标正常或仅轻微偏离，通常随访即可
   - amber: 有1-2项明确异常或存在需要复查/结合症状判断的问题
   - red: 有多项明显异常、关键器官相关异常、或提示需尽快就医的情况
8. key_findings 中每项都必须包含 name/value/reference/status/risk/explanation，
   其中 status 仅允许 normal/high/low，risk 仅允许 green/amber/red。
9. disclaimer 必须原样输出：{REPORT_DISCLAIMER}

特别关注：{report_focus}

严格输出JSON对象，不要输出Markdown，不要输出代码块，不要补充JSON外文字。
JSON格式如下：
{{
  "report_type": "blood_routine" | "liver_function" | "thyroid" | "immunology" | "comprehensive" | "general",
  "risk_level": "green" | "amber" | "red",
  "summary": "200字以内白友易懂总结",
  "key_findings": [
    {{
      "name": "指标名称",
      "value": "检测值",
      "reference": "参考范围",
      "status": "normal" | "high" | "low",
      "risk": "green" | "amber" | "red",
      "explanation": "通俗解释"
    }}
  ],
  "recommendations": ["建议1", "建议2"],
  "disclaimer": "{REPORT_DISCLAIMER}"
}}"""


def _default_report_interpretation(
    report_type: str,
    summary: str,
    risk_level: str = "amber",
    recommendations: Optional[List[str]] = None,
) -> dict:
    return {
        "report_type": report_type,
        "risk_level": risk_level,
        "summary": summary,
        "key_findings": [],
        "recommendations": recommendations
        or ["若报告存在异常箭头或超出参考范围，请咨询医生进一步判断。"],
        "disclaimer": REPORT_DISCLAIMER,
    }


def _normalize_report_type(value: object, fallback: str) -> str:
    allowed = {
        "blood_routine",
        "liver_function",
        "thyroid",
        "immunology",
        "comprehensive",
        "general",
    }
    if isinstance(value, str) and value in allowed:
        return value
    return fallback


def _normalize_level(value: object, allowed: set[str], fallback: str) -> str:
    if isinstance(value, str) and value in allowed:
        return value
    return fallback


def _normalize_key_findings(value: object) -> List[dict]:
    findings: List[dict] = []
    if not isinstance(value, list):
        return findings

    for item in value:
        if not isinstance(item, dict):
            continue

        name = str(item.get("name", "")).strip()
        if not name:
            continue

        findings.append(
            {
                "name": name,
                "value": str(item.get("value", "未提供")).strip() or "未提供",
                "reference": str(item.get("reference", "未提供")).strip() or "未提供",
                "status": _normalize_level(
                    item.get("status"), {"normal", "high", "low"}, "normal"
                ),
                "risk": _normalize_level(
                    item.get("risk"), {"green", "amber", "red"}, "green"
                ),
                "explanation": str(item.get("explanation", "")).strip()
                or "建议结合医生意见综合判断。",
            }
        )

    return findings


def _normalize_recommendations(value: object) -> List[str]:
    if not isinstance(value, list):
        return ["若报告存在异常箭头或超出参考范围，请咨询医生进一步判断。"]

    recommendations = [str(item).strip() for item in value if str(item).strip()]
    return recommendations or [
        "若报告存在异常箭头或超出参考范围，请咨询医生进一步判断。"
    ]


def _interpret_report(doc_text: str, question: str) -> dict:
    """Use LLM to interpret a medical report."""
    config = get_llm_config()
    report_type = _detect_report_type(doc_text)
    if config["provider"] == "none":
        return _default_report_interpretation(
            report_type=report_type,
            risk_level="amber",
            summary="AI服务未配置，暂时无法自动解读报告。建议结合原始报告中的异常箭头、参考范围和医生意见综合判断。",
        )

    client = openai.OpenAI(api_key=config["api_key"], base_url=config["base_url"])
    prompt = f"""请基于以下内容解读报告。

报告类型（供参考）: {REPORT_TYPE_LABELS.get(report_type, "综合医疗报告")} ({report_type})
用户问题: {question}

报告原文：
{doc_text[:3000]}
"""

    try:
        response = client.chat.completions.create(
            model=config["chat_model"],
            messages=[
                {"role": "system", "content": _build_report_prompt(report_type)},
                {"role": "user", "content": prompt},
            ],
            temperature=0.1,
        )
        content = (response.choices[0].message.content or "").strip()
        if content.startswith("```"):
            content = content.strip("`")
            if content.startswith("json"):
                content = content[4:].strip()
        result = json.loads(content)
        if not isinstance(result, dict):
            raise ValueError("report interpretation is not an object")

        normalized_report_type = _normalize_report_type(
            result.get("report_type"), report_type
        )
        return {
            "report_type": normalized_report_type,
            "risk_level": _normalize_level(
                result.get("risk_level"), {"green", "amber", "red"}, "amber"
            ),
            "summary": str(result.get("summary", "报告已解读")).strip() or "报告已解读",
            "key_findings": _normalize_key_findings(result.get("key_findings")),
            "recommendations": _normalize_recommendations(
                result.get("recommendations")
            ),
            "disclaimer": str(result.get("disclaimer", REPORT_DISCLAIMER)).strip()
            or REPORT_DISCLAIMER,
        }
    except Exception as exc:
        logger.warning("Report interpretation failed: %s", str(exc))
        return _default_report_interpretation(
            report_type=report_type,
            risk_level="amber",
            summary="报告解读失败，建议查看原始报告中的异常提示，并尽快咨询医生或检验科进一步说明。",
        )


def _generate_diary_draft(question: str, answer: str, action_cards: list) -> dict:
    """Generate a diary draft from the conversation and action cards."""
    config = get_llm_config()
    today = date.today().isoformat()

    if config["provider"] == "none":
        return {
            "type": "diary",
            "title": f"{today} 病情日记",
            "content": f"今日问题: {question}\n\nAI回答摘要: {answer[:200]}",
            "date": today,
            "privacy": "private",
            "saved": False,
        }

    client = openai.OpenAI(api_key=config["api_key"], base_url=config["base_url"])

    cards_context = ""
    for card in action_cards:
        if card.get("type") == "vasi":
            cards_context += (
                f"\nVASI评估: 评分{card.get('vasiScore', 0)}分, "
                f"部位{card.get('bodySite', '其他')}, 分期{card.get('stage', '未知')}"
            )
        elif card.get("type") == "report":
            cards_context += f"\n体检报告解读: {card.get('summary', '')}"

    prompt = f"""基于以下内容，帮用户生成一篇白癜风日记草稿。要求：
1. 语言温暖、鼓励，第一人称
2. 包含今日的问题和AI的建议
3. 包含评估/报告的关键信息
4. 200字以内

用户问题: {question}
AI回答: {answer[:500]}
{cards_context}

请直接输出日记内容（HTML格式），不要加标题。"""

    try:
        response = client.chat.completions.create(
            model=config["chat_model"],
            messages=[{"role": "user", "content": prompt}],
            temperature=0.5,
        )
        content = (response.choices[0].message.content or "").strip()
    except Exception:
        content = f"<p>今日咨询了关于{question[:30]}的问题。{answer[:100]}</p>"

    return {
        "type": "diary",
        "title": f"{today} 病情日记",
        "content": content,
        "date": today,
        "privacy": "private",
        "saved": False,
    }


def generate_answer(
    query: str,
    docs: List[Document],
    conversation_history: Optional[List[dict]] = None,
    mode: Optional[str] = None,
) -> str:
    config = get_llm_config()

    if config["provider"] == "none":
        return "抱歉，AI 问答服务暂未配置，请管理员设置 DASHSCOPE_API_KEY。"

    client = openai.OpenAI(
        api_key=config["api_key"],
        base_url=config["base_url"],
    )

    messages = _build_llm_messages(query, docs, conversation_history, mode)
    temperature = 0.7 if mode == "counseling" else 0.3

    logger.info(
        "调用 chat API: provider=%s, model=%s, mode=%s, temperature=%s",
        config["provider"],
        config["chat_model"],
        mode or "knowledge",
        temperature,
    )

    try:
        response = client.chat.completions.create(
            model=config["chat_model"],
            messages=messages,
            temperature=temperature,
        )
        return response.choices[0].message.content.strip()
    except openai.APIError as e:
        logger.error("Chat API 调用失败: %s", str(e))
        return f"AI 问答服务暂时不可用，请稍后再试。（错误: {type(e).__name__}）"


def answer_question_stream(
    query: str,
    docs: List[Document],
    conversation_history: Optional[List[dict]] = None,
    has_attachments: bool = False,
    mode: Optional[str] = None,
):
    """Stream answer tokens one by one, yielding each token as it arrives."""
    config = get_llm_config()

    if config["provider"] == "none":
        yield "抱歉，AI 问答服务暂未配置，请管理员设置 DASHSCOPE_API_KEY。"
        return

    client = openai.OpenAI(
        api_key=config["api_key"],
        base_url=config["base_url"],
    )

    messages = _build_llm_messages(query, docs, conversation_history, mode)
    temperature = 0.7 if mode == "counseling" else 0.3

    if has_attachments:
        attachment_note = "\n注意：用户本次对话已上传了附件（图片或体检报告），你已经在解读结果了，不要在回答中再推荐上传方式或引导用户去上传报告。"
        messages[0]["content"] += attachment_note

    logger.info(
        "调用 chat API (stream): provider=%s, model=%s, mode=%s, temperature=%s",
        config["provider"],
        config["chat_model"],
        mode or "knowledge",
        temperature,
    )

    try:
        response = client.chat.completions.create(
            model=config["chat_model"],
            messages=messages,
            temperature=temperature,
            stream=True,
        )
        for chunk in response:
            delta = chunk.choices[0].delta
            if delta and delta.content:
                yield delta.content
    except openai.APIError as e:
        logger.error("Chat API streaming 调用失败: %s", str(e))
        yield f"AI 问答服务暂时不可用，请稍后再试。（错误: {type(e).__name__}）"


def answer_question(
    db: Session,
    question: str,
    conversation_id: str = None,
    user_id: int = None,
    mode: Optional[str] = None,
) -> QuestionResponse:
    """完整的RAG问答流程"""
    results = search_documents(db, question)
    docs = [doc for doc, score in results if score > 0.5]

    if not docs:
        docs = [doc for doc, score in results[:3]]

    sources = [
        Source(
            title=doc.title,
            url=doc.source_url if doc.source_url else "",
            snippet=doc.content[:200] + "..."
            if len(doc.content) > 200
            else doc.content,
            source_tier=doc.source_tier or "C",
            source_name=doc.source or "",
            authority_weight=float(doc.authority_weight or 1.0),
        )
        for doc in docs
    ]

    conversation_history = None
    if conversation_id:
        history = (
            db.query(Message)
            .filter(Message.conversation_id == conversation_id)
            .order_by(Message.created_at)
            .all()
        )
        conversation_history = [
            {"role": msg.role, "content": msg.content} for msg in history
        ]

        conv = (
            db.query(Conversation)
            .filter(Conversation.conversation_id == conversation_id)
            .first()
        )
        if not conv:
            conv = Conversation(conversation_id=conversation_id, user_id=user_id)
            db.add(conv)
            db.commit()

    answer = generate_answer(question, docs, conversation_history, mode)

    if conversation_id:
        user_msg = Message(
            conversation_id=conversation_id, role="user", content=question
        )
        db.add(user_msg)
        assistant_msg = Message(
            conversation_id=conversation_id, role="assistant", content=answer
        )
        db.add(assistant_msg)
        db.commit()

    return QuestionResponse(answer=answer, sources=sources)


def add_document(
    db: Session,
    title: str,
    content: str,
    source: str = None,
    source_url: str = None,
    category: str = None,
    source_tier: str = "C",
    authority_weight: float = 1.0,
    pub_date: str = None,
    compute_embedding: bool = False,
) -> Document:
    """添加文档到知识库。默认不计算 embedding（每月1号批量增量向量化），如需即时计算可设 compute_embedding=True。"""
    import json

    embedding_value = None
    if compute_embedding:
        content_truncated = content[:8000]
        embedding = get_embedding(content_truncated)
        embedding_value = json.dumps(embedding)

    doc = Document(
        title=title,
        content=content,
        source=source,
        source_url=source_url,
        category=category,
        source_tier=source_tier,
        authority_weight=authority_weight,
        pub_date=pub_date,
        embedding=embedding_value,
    )
    db.add(doc)
    db.commit()
    db.refresh(doc)
    return doc


def batch_embed_unembedded(db: Session) -> dict:
    """对 embedding=NULL 的文档做增量向量化（每月1号批量执行）。

    Returns:
        dict with embedded_count, failed_count, total, skipped_count
    """
    import json
    import time

    from web.backend.database.models import Document

    unembedded = db.query(Document).filter(Document.embedding.is_(None)).all()
    total = len(unembedded)

    if total == 0:
        logger.info("没有需要向量化的新文档")
        return {"embedded_count": 0, "failed_count": 0, "total": 0, "skipped_count": 0}

    logger.info("开始增量向量化: %d 篇文档需要处理", total)

    embedded_count = 0
    failed_count = 0

    for i, doc in enumerate(unembedded):
        try:
            content_truncated = doc.content[:8000] if doc.content else ""
            if not content_truncated.strip():
                failed_count += 1
                logger.warning("文档 %d 内容为空，跳过", doc.id)
                continue

            embedding = get_embedding(content_truncated)
            doc.embedding = json.dumps(embedding)
            db.commit()

            embedded_count += 1
            if (embedded_count) % 10 == 0:
                logger.info(
                    "向量化进度: %d/%d 已完成 (%.1f%%)",
                    embedded_count,
                    total,
                    embedded_count / total * 100,
                )

            # 限速：每次 API 调用间隔 1.5 秒，避免触发 rate limit
            time.sleep(1.5)

        except Exception as e:
            failed_count += 1
            logger.error("文档 %d 向量化失败: %s", doc.id, str(e))
            db.rollback()
            continue

    logger.info(
        "增量向量化完成: 成功 %d, 失败 %d, 总计 %d",
        embedded_count,
        failed_count,
        total,
    )

    return {
        "embedded_count": embedded_count,
        "failed_count": failed_count,
        "total": total,
        "skipped_count": 0,
    }
